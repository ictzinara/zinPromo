from os.path import isfile
import xlrd
import logging
import openpyxl
from django.db.models import Q
from django.shortcuts import render
from django.http import HttpResponse
from django.template import loader
from django.core.exceptions import ObjectDoesNotExist
from Api.models import Province, PromotionWeeklyDraw, LicenseDb
from zinPromo.forms import PromotionApplicantForm ,AddVehicleForm
from zinPromo.tables import WeeklyDrawTable

# import pandas lib as pd
import pandas as pd

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


# if isfile(filepath):
#     data_reader = csv.reader(open(filepath, 'rb'), delimiter='|')
#     logger.info("importing to %s" % model)
#     # vaciar la tabla
#     model.objects.all().delete()
#
#     if type(formatter) == list:
#         format = lambda d: dict(zip(formatter, d))
#     else:
#         format = formatter
#     i = 0
#     for row_value_list in data_reader:
#         data = format(row_value_list)
#         if data:
#             obj = model(**data)
#             obj.save()
#             i = i + 1
#         else:
#             logger.error("incomplete data %s" % data)
#
#     logger.info("%s objects imported to %s" % (i, model))
# else:
#     logger.error("temp import file not found")
def index(request):
    template = loader.get_template('index.html')
    provinces = Province.objects.all()
    context = {
        'data': [],
        'provinces': provinces,
    }
    return HttpResponse(template.render(context, request))


def upload(request):
    if "GET" == request.method:
        return render(request, 'upload.html', {})
    else:
        excel_file = request.FILES["excel_file"]
        # you may put validations here to check extension or file size
        wb = openpyxl.load_workbook(excel_file)
        # getting a particular sheet by name out of many sheets
        worksheet = wb["Sheet1"]
        # print(worksheet)
        # read by default 1st sheet of an excel file
        dataframe1 = pd.read_excel(excel_file)
        dataframe1 = dataframe1.rename(columns={
            'REGISTRATION NO': 'REGISTRATION_NO',
            'LICENSE STATUS': 'LICENSE_STATUS',
            'LAST LICENSING TRANSACTION': 'LAST_LICENSING_TRANSACTION',
            'PENALTY AMOUNT': 'PENALTY_AMOUNT',
            'ARREAR AMOUNT': 'ARREAR_AMOUNT',
            'LICENCE EXPIRY DATE': 'LICENCE_EXPIRY_DATE',
        })
        print('Data Frame')
        # print(dataframe1)
        i = 0
        for row in dataframe1.itertuples():
            data_row = {
                "EXPIRED": '-',
                "REGISTRATION_NO": row.REGISTRATION_NO,
                "LICENSE_STATUS": row.LICENSE_STATUS,
                "LAST_LICENSING_TRANSACTION": row.LAST_LICENSING_TRANSACTION,
                "PENALTY_AMOUNT": row.PENALTY_AMOUNT,
                "ARREAR_AMOUNT": row.ARREAR_AMOUNT,
                "BLACKLISTED": row.BLACKLISTED,
                "LICENCE_EXPIRY_DATE": row.LICENCE_EXPIRY_DATE,
            }

            db_record = LicenseDb.objects.filter(REGISTRATION_NO=row.REGISTRATION_NO)
            if db_record:
                print("record Exists")
            else:
                form = AddVehicleForm(data=data_row)
                if form.is_valid():
                    if form.save():
                        i = i+1
                        print("New Record Created")
                    else:
                        logger.error("Error Saving data %s" % data_row)

            print(data_row)
            # print(data)
        logger.info("%s objects imported to %s" % (i, LicenseDb))

        excel_data = list()
        # iterating over the rows and
        # getting value from each cell in row
        for row in worksheet.iter_rows():
            row_data = list()
            # print(row)
            for cell in row:
                row_data.append(str(cell.value))
            excel_data.append(row_data)
        return render(request, 'upload.html', {"excel_data": excel_data})


def provincial_winners(request):
    template = loader.get_template('winners.html')
    provinces = Province.objects.all()
    provincial_winners = PromotionWeeklyDraw.objects.all()
    table_ww = WeeklyDrawTable()
    context = {
        'data': [],
        'table_ww': table_ww,
        'provinces': provinces,
        'provincial_winners': provincial_winners,
    }
    return HttpResponse(template.render(context, request))


def vrn_eligibility(request):
    if request.POST:
        v_reg = request.POST.get('vrn').upper()
        prvn = request.POST.get('prvn')
        try:
            mydata = LicenseDb.objects.get(REGISTRATION_NO=v_reg)
            stat = "Licenced"
        except ObjectDoesNotExist:
            stat = "UnLicenced"
            mydata = []
        # mydata = vlicDb.objects.filter(Q(Regno=v_reg) | Q(firstname='Tobias')).values()
    template = loader.get_template('apply.html')
    context = {
        'data': [],
        'v_reg': v_reg,
        'prvn': prvn,
        'status': stat,
        'reg_validity': mydata,
    }
    return HttpResponse(template.render(context, request))


def apply(request):
    if request.POST:
        form = PromotionApplicantForm(data=request.POST)
        if form.is_valid():
            form.save()
            return HttpResponse("Data Saved Successfully")
        else:
            return HttpResponse("Form Failed to Validate")
    template = loader.get_template('tours.html')
    context = {
        'data': [],
    }
    return HttpResponse(template.render(context, request))


def import_xls_to_db(filepath, model, model_mapping):
    """ Imports data from .xls to db"""
    first_row = 1

    if isfile(filepath):
        book = xlrd.open_workbook(filepath, encoding_override='cp1252')
        sheet = book.sheet_by_index(0)
        nr_rows = sheet.nrows
        for row in range(first_row, nr_rows):
            row_value_list = [cell.value for cell in sheet.row(row)]
            if (type(model_mapping) == dict):
                da = zip(model_mapping.keys(), row_value_list)
                data = dict([(i[0], model_mapping[i[0]](i[1])) for i in da])
            else:
                data = dict(zip(model_mapping, row_value_list))
            print(data)
            obj = model(**data)
            obj.save()
            msg = "imported %s" % data
            logger.info(msg)
